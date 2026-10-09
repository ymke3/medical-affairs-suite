import os
import io
import streamlit as st
import requests
import xml.etree.ElementTree as ET
from dotenv import load_dotenv
from docx import Document
from google import genai

load_dotenv(override=True)

# API Key ophalen via Secrets of .env
gemini_key = None
if "GEMINI_API_KEY" in st.secrets:
    gemini_key = str(st.secrets["GEMINI_API_KEY"]).strip("\"' ")
elif os.getenv("GEMINI_API_KEY"):
    gemini_key = str(os.getenv("GEMINI_API_KEY")).strip("\"' ")

st.set_page_config(
    page_title="Medical Affairs Enterprise Suite",
    page_icon="🩺",
    layout="wide"
)

if not gemini_key:
    st.error("⚠️ Geen geldige GEMINI_API_KEY gevonden in Streamlit Secrets. Voeg deze toe via Settings -> Secrets.")
    st.stop()

@st.cache_resource
def get_gemini_client(api_key):
    return genai.Client(api_key=api_key)

client = get_gemini_client(gemini_key)

# --- Functies voor Data en Export ---

def fetch_pubmed_data(query, max_results=5):
    esearch_url = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esearch.fcgi"
    search_params = {
        "db": "pubmed",
        "term": query,
        "retmode": "json",
        "retmax": max_results,
        "sort": "pub_date"
    }
    
    try:
        search_resp = requests.get(esearch_url, params=search_params, timeout=12)
        id_list = search_resp.json().get("esearchresult", {}).get("idlist", [])
        
        if not id_list:
            return []
            
        efetch_url = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi"
        fetch_params = {
            "db": "pubmed",
            "id": ",".join(id_list),
            "retmode": "xml"
        }
        fetch_resp = requests.get(efetch_url, params=fetch_params, timeout=12)
        root = ET.fromstring(fetch_resp.content)
        
        studies = []
        for article in root.findall(".//PubmedArticle"):
            pmid = article.findtext(".//MedlineCitation/PMID") or "Onbekend"
            title = article.findtext(".//ArticleTitle") or "Geen titel"
            abstract_texts = article.findall(".//Abstract/AbstractText")
            abstract = " ".join([elem.text for elem in abstract_texts if elem.text]) or "Geen abstract beschikbaar."
            journal = article.findtext(".//Journal/Title") or "Tijdschrift onbekend"
            pub_year = article.findtext(".//JournalIssue/PubDate/Year") or "Onbekend"
            
            studies.append({
                "pmid": pmid,
                "title": title,
                "abstract": abstract,
                "source": f"{journal} ({pub_year})"
            })
        return studies
    except Exception as e:
        st.error(f"Fout bij benaderen PubMed: {e}")
        return []

def create_docx(content_title, text_content):
    doc = Document()
    doc.add_heading(content_title, level=0)
    for paragraph in text_content.split("\n\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io

# --- Interface Setup ---

st.title("🩺 Medical Affairs & Clinical Intelligence Suite")
st.caption("Gespecialiseerd platform voor Medical Directors, MSL teams en Clinical Research Consultants")

tab_briefing, tab_adboard, tab_kol = st.tabs([
    "📑 Executive Briefings",
    "📝 Advisory Boards & Meetings",
    "👨‍⚕️ KOL & Investigator Profiler"
])

# --- TAB 1: EXECUTIVE BRIEFINGS ---
with tab_briefing:
    st.subheader("Weekly Clinical & Competitor Briefing")
    st.markdown("Genereer een direct presenteerbare samenvatting van de nieuwste klinische data en concurrentie-ontwikkelingen.")
    
    col1, col2, col3 = st.columns([3, 1, 1])
    with col1:
        query_input = st.text_input("Zoekopdracht (Indicatie / Target / Compound):", "NSCLC AND immunotherapy AND Phase 3")
    with col2:
        studies_limit = st.selectbox("Aantal publicaties:", [3, 5, 8], index=1)
    with col3:
        target_audience = st.selectbox("Focus:", ["Medical Affairs & Strategie", "Klinische Methodologie & Safety", "Concurrentie & Pijplijn"])

    if st.button("Genereer Intelligence Briefing", type="primary"):
        with st.spinner("PubMed raadplegen en klinische synthese formuleren..."):
            fetched_studies = fetch_pubmed_data(query_input, max_results=studies_limit)
            
            if not fetched_studies:
                st.warning("Geen publicaties gevonden voor deze criteria.")
            else:
                studies_formatted = "\n\n".join([
                    f"PMID: {s['pmid']}\nTitel: {s['title']}\nBron: {s['source']}\nAbstract: {s['abstract']}"
                    for s in fetched_studies
                ])
                
                briefing_prompt = f"""
                Je bent een Senior Director Medical Affairs met decennia aan ervaring in farma en biotech.
                Schrijf op basis van de onderstaande klinische studies een professionele, diepgaande Executive Briefing in het Nederlands.
                Focus: {target_audience}

                Data:
                {studies_formatted}

                Hanteer een strakke, professionele structuur:
                1. EXECUTIVE SUMMARY: 3 strategische kernconclusies voor het medische en commerciële team.
                2. KLINISCHE EVALUATIE (Per studie):
                   - Opzet & Patiëntpopulatie (inclusief fase, steekproefgrootte N)
                   - Primaire & Secundaire Uitkomsten (exacte statistiek benoemen: HR, p-waarden, OS/PFS, toxiciteit)
                   - Klinische & Strategische Duiding (wat betekent dit t.o.v. de huidige zorgstandaard?)
                3. COMPETITIVE & PIPELINE IMPACT: Welke invloed heeft dit op lopende studies of concurrerende geneesmiddelen?
                4. ADVISORY BOARD QUESTIONS: 2 scherpe, methodologische discussievragen voor medisch specialisten.
                """
                
                try:
                    resp = client.models.generate_content(
                        model="gemini-3.8-flash",
                        contents=briefing_prompt
                    )
                    st.markdown(resp.text)
                    
                    docx_file = create_docx(f"Medical Briefing - {query_input}", resp.text)
                    st.download_button(
                        label="📥 Download als Word Document (.docx)",
                        data=docx_file,
                        file_name="Medical_Briefing.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
                except Exception as api_err:
                    st.error(f"Foutdetails van Google API: {api_err}")

# --- TAB 2: ADVISORY BOARDS & MEETINGS ---
with tab_adboard:
    st.subheader("Meeting & Advisory Board Synthesizer")
    st.markdown("Zet ruwe transcripties of losse vergadernotities om in gestructureerde farma-notulen en actieplannen.")
    
    meeting_context = st.text_input("Onderwerp / Type Overleg:", "National Advisory Board Oncologie")
    raw_input = st.text_area("Plak ruwe gespreksnotities of transcript:", height=220, placeholder="Arts A benoemde zorgen over graad 3 trombocytopenie... Besproken dat protocolamendement nodig is voor cohort B...")
    
    if st.button("Synthetiseer Notulen"):
        if not raw_input.strip():
            st.warning("Voer eerst notities in.")
        else:
            with st.spinner("Medische discussies structureren..."):
                notes_prompt = f"""
                Je bent verantwoordelijk voor Medical Governance & Reporting bij een toonaangevend farmabedrijf.
                Zet de onderstaande ruwe notities van een '{meeting_context}' om in een formele executive samenvatting volgens farma-standaarden.

                Ruwe tekst:
                {raw_input}

                Structuur:
                - 🎯 Medisch-Wetenschappelijke Inzichten (wat waren de belangrijkste meningen en klinische feedback?)
                - ⚖️ Strategische Besluiten (welke consensus of formele beslissingen zijn genomen?)
                - 📋 Actielijst & Verantwoordelijkheden (duidelijke tabel met: Actiepunt | Eigenaar | Prioriteit)
                """
                
                try:
                    resp_notes = client.models.generate_content(
                        model="gemini-3.8-flash",
                        contents=notes_prompt
                    )
                    st.markdown(resp_notes.text)
                    docx_notes = create_docx(f"Notulen - {meeting_context}", resp_notes.text)
                    st.download_button(
                        label="📥 Download Notulen (.docx)",
                        data=docx_notes,
                        file_name="Meeting_Minutes.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
                except Exception as api_err:
                    st.error(f"Foutdetails van Google API: {api_err}")

# --- TAB 3: KOL PROFILER ---
with tab_kol:
    st.subheader("Key Opinion Leader (KOL) & Expert Profiler")
    st.markdown("Analyseer de wetenschappelijke footprint van een medisch specialist of hoofdonderzoeker.")
    
    col_k1, col_k2 = st.columns([3, 1])
    with col_k1:
        investigator_name = st.text_input("Auteur / Specialist (bijv. 'van Dongen J' of achternaam met initiaal):", "Peters S")
    with col_k2:
        article_count = st.selectbox("Max artikelen:", [5, 10, 15], index=1)
        
    if st.button("Genereer KOL Dossier"):
        with st.spinner(f"Publicatiehistorie en netwerk analyseren voor {investigator_name}..."):
            kol_articles = fetch_pubmed_data(f"{investigator_name}[Author]", max_results=article_count)
            
            if not kol_articles:
                st.warning("Geen recente publicaties gevonden voor deze specialist op PubMed.")
            else:
                kol_records = "\n\n".join([f"- {a['title']} | {a['source']} (PMID: {a['pmid']})" for a in kol_articles])
                
                kol_prompt = f"""
                Analyseer de wetenschappelijke publicatielijst van expert '{investigator_name}' voor een Medical Affairs afdeling:

                Publicaties:
                {kol_records}

                Stel een compact executive profiel op:
                1. Primair Onderzoeks- en Expertisedomein (waar staat deze onderzoeker om bekend?)
                2. Recente Speerpunten & Trends (welke moleculen, pathways of studiefases behandelen de nieuwste papers?)
                3. Geschiktheid voor Samenwerking (bijv. als Advisory Board lid, Steering Committee participant, of MSL engagement)
                """
                
                try:
                    resp_kol = client.models.generate_content(
                        model="gemini-3.8-flash",
                        contents=kol_prompt
                    )
                    st.markdown(resp_kol.text)
                    docx_kol = create_docx(f"KOL Profiel - {investigator_name}", resp_kol.text)
                    st.download_button(
                        label="📥 Download Dossier (.docx)",
                        data=docx_kol,
                        file_name=f"KOL_Profiel_{investigator_name.replace(' ', '_')}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
                except Exception as api_err:
                    st.error(f"Foutdetails van Google API: {api_err}")
